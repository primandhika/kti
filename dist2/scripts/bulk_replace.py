#!/usr/bin/env python3
"""
Bulk Find & Replace in .docx using CSV mapping.

CSV format expected:
  "no","before(verbatim)","after(sugesti)"

The script handles text split across multiple Word runs by
reconstructing paragraph text and carefully replacing across run boundaries.
"""

import csv
import copy
import sys
import os
from docx import Document


def replace_in_paragraph(paragraph, old_text, new_text):
    """
    Replace old_text with new_text in a paragraph, handling text
    that may be split across multiple runs.
    
    Returns True if a replacement was made, False otherwise.
    """
    # First check if the full paragraph text contains old_text
    full_text = paragraph.text
    if old_text not in full_text:
        return False

    # Strategy: reconstruct runs to perform the replacement
    # We need to handle the case where old_text spans multiple runs
    
    # Build a mapping: for each character position in full_text,
    # record which run and position within that run it belongs to
    char_map = []  # list of (run_index, char_index_in_run)
    for run_idx, run in enumerate(paragraph.runs):
        for char_idx in range(len(run.text)):
            char_map.append((run_idx, char_idx))
    
    if not char_map:
        return False
    
    # Find the position of old_text in full_text
    start_pos = full_text.find(old_text)
    if start_pos == -1:
        return False
    
    end_pos = start_pos + len(old_text)
    
    # Determine which runs are affected
    start_run_idx, start_char_idx = char_map[start_pos]
    end_run_idx, end_char_idx = char_map[end_pos - 1]
    
    # Get the formatting from the first run of the match
    first_run = paragraph.runs[start_run_idx]
    
    if start_run_idx == end_run_idx:
        # Simple case: old_text is entirely within one run
        run = paragraph.runs[start_run_idx]
        run.text = run.text[:start_char_idx] + new_text + run.text[end_char_idx + 1:]
    else:
        # Complex case: old_text spans multiple runs
        # 1. Put new_text in the first affected run (keeping prefix)
        first_run = paragraph.runs[start_run_idx]
        prefix = first_run.text[:start_char_idx]
        first_run.text = prefix + new_text
        
        # 2. Clear intermediate runs
        for run_idx in range(start_run_idx + 1, end_run_idx):
            paragraph.runs[run_idx].text = ""
        
        # 3. Trim the last affected run (keeping suffix)
        last_run = paragraph.runs[end_run_idx]
        last_run.text = last_run.text[end_char_idx + 1:]
    
    return True


def replace_in_table(table, old_text, new_text):
    """Replace text in all cells of a table."""
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                while old_text in paragraph.text:
                    if replace_in_paragraph(paragraph, old_text, new_text):
                        count += 1
                    else:
                        break
    return count


def replace_in_document(doc, old_text, new_text):
    """Replace text throughout the entire document."""
    count = 0
    
    # Replace in main body paragraphs
    for paragraph in doc.paragraphs:
        while old_text in paragraph.text:
            if replace_in_paragraph(paragraph, old_text, new_text):
                count += 1
            else:
                break
    
    # Replace in tables
    for table in doc.tables:
        count += replace_in_table(table, old_text, new_text)
    
    # Replace in headers and footers
    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header and header.is_linked_to_previous is False:
                for paragraph in header.paragraphs:
                    while old_text in paragraph.text:
                        if replace_in_paragraph(paragraph, old_text, new_text):
                            count += 1
                        else:
                            break
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer and footer.is_linked_to_previous is False:
                for paragraph in footer.paragraphs:
                    while old_text in paragraph.text:
                        if replace_in_paragraph(paragraph, old_text, new_text):
                            count += 1
                        else:
                            break
    
    return count


def load_replacements(csv_path):
    """Load before/after pairs from CSV file."""
    replacements = []
    with open(csv_path, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            before = row['before(verbatim)'].strip()
            after = row['after(sugesti)'].strip()
            if before and after and before != after:
                replacements.append((before, after))
    return replacements


def main():
    docx_path = '/home/primandhika/artikel/dist2/[Hasil] DRAF7 - Disertasi Restu.docx'
    csv_path = '/home/primandhika/artikel/dist2/Z_before_after_revisi_DRAF7.csv'
    
    # Output path
    output_dir = '/home/primandhika/artikel/dist2/outputs'
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, '[Hasil] DRAF7 - Disertasi Restu - REVISED.docx')
    
    print("=" * 60)
    print("BULK FIND & REPLACE - DOCX")
    print("=" * 60)
    
    # Load replacements from CSV
    print(f"\n📄 Loading replacements from CSV...")
    replacements = load_replacements(csv_path)
    print(f"   Found {len(replacements)} replacement pairs")
    
    # Load the document
    print(f"\n📂 Loading document...")
    doc = Document(docx_path)
    print(f"   Document loaded successfully")
    print(f"   Paragraphs: {len(doc.paragraphs)}")
    print(f"   Tables: {len(doc.tables)}")
    
    # Perform replacements
    print(f"\n🔄 Performing replacements...")
    total_replacements = 0
    results = []
    
    for i, (before, after) in enumerate(replacements, 1):
        count = replace_in_document(doc, before, after)
        total_replacements += count
        status = "✅" if count > 0 else "⚠️  NOT FOUND"
        short_before = before[:80] + "..." if len(before) > 80 else before
        results.append((i, count, status, short_before))
        print(f"   [{i:2d}/{len(replacements)}] {status} (count={count}) | {short_before}")
    
    # Save the modified document
    print(f"\n💾 Saving modified document...")
    doc.save(output_path)
    print(f"   Saved to: {output_path}")
    
    # Summary
    print(f"\n{'=' * 60}")
    print(f"SUMMARY")
    print(f"{'=' * 60}")
    print(f"   Total replacement pairs:  {len(replacements)}")
    print(f"   Successful replacements:  {sum(1 for _, c, _, _ in results if c > 0)}")
    print(f"   Not found:               {sum(1 for _, c, _, _ in results if c == 0)}")
    print(f"   Total text replacements:  {total_replacements}")
    print(f"   Output file:             {output_path}")
    print(f"{'=' * 60}")


if __name__ == '__main__':
    main()
