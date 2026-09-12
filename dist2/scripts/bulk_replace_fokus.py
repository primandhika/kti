#!/usr/bin/env python3
"""
Round 2b - Fokus Penelitian corrections.

The CSV 'before' texts were already replaced by Round 1.
So we need to replace the Round 1 'after' texts with the new corrections.

Mapping (using the actual text currently in the document):
1. "empat subfokus" paragraph -> change to "tiga subfokus" version
2. "Karakteristik dan kebutuhan belajar..." -> new subfokus 1
3. "Rancangan media web microlearning..." -> new subfokus 2
4. "Proses pengembangan dan implementasi..." -> new subfokus 3
5. "Kelayakan dan efektivitas..." -> DELETE (was subfokus 4, now removed)
"""

import os
from docx import Document


def replace_in_paragraph(paragraph, old_text, new_text):
    full_text = paragraph.text
    if old_text not in full_text:
        return False
    char_map = []
    for run_idx, run in enumerate(paragraph.runs):
        for char_idx in range(len(run.text)):
            char_map.append((run_idx, char_idx))
    if not char_map:
        return False
    start_pos = full_text.find(old_text)
    if start_pos == -1:
        return False
    end_pos = start_pos + len(old_text)
    start_run_idx, start_char_idx = char_map[start_pos]
    end_run_idx, end_char_idx = char_map[end_pos - 1]
    if start_run_idx == end_run_idx:
        run = paragraph.runs[start_run_idx]
        run.text = run.text[:start_char_idx] + new_text + run.text[end_char_idx + 1:]
    else:
        first_run = paragraph.runs[start_run_idx]
        prefix = first_run.text[:start_char_idx]
        first_run.text = prefix + new_text
        for run_idx in range(start_run_idx + 1, end_run_idx):
            paragraph.runs[run_idx].text = ""
        last_run = paragraph.runs[end_run_idx]
        last_run.text = last_run.text[end_char_idx + 1:]
    return True


def remove_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)


def main():
    docx_path = '/home/primandhika/artikel/dist2/outputs/[Hasil] DRAF7 - Disertasi Restu - REVISED.docx'

    # These are the texts currently in the document (from Round 1 replacements)
    # mapped to the new corrections from the fokus penelitian CSV
    replacements = [
        # 1. Change "empat subfokus" to "tiga subfokus"
        (
            "Penelitian ini berfokus pada pengembangan dan evaluasi media web microlearning berbasis Teknik Feynman untuk meningkatkan kemampuan metakognitif dalam keterampilan berbicara mahasiswa Pendidikan Bahasa dan Sastra Indonesia. Fokus tersebut dijabarkan ke dalam empat subfokus penelitian:",
            "Penelitian ini berfokus pada pengembangan dan evaluasi media web microlearning berbasis Teknik Feynman untuk meningkatkan kemampuan metakognitif dalam keterampilan berbicara mahasiswa Pendidikan Bahasa dan Sastra Indonesia. Fokus tersebut dijabarkan ke dalam tiga subfokus penelitian:",
        ),
        # 2. Subfokus 1 -> new text
        (
            "Karakteristik dan kebutuhan belajar mahasiswa dalam pembelajaran keterampilan berbicara berbasis metakognitif.",
            "Pengembangan media web microlearning berbasis Teknik Feynman berdasarkan karakteristik dan kebutuhan belajar mahasiswa dengan menggunakan model ASSURE.",
        ),
        # 3. Subfokus 2 -> new text
        (
            "Rancangan media web microlearning berbasis Teknik Feynman berdasarkan karakteristik mahasiswa dan tujuan pembelajaran.",
            "Kelayakan media web microlearning berbasis Teknik Feynman berdasarkan penilaian ahli dan respons pengguna.",
        ),
        # 4. Subfokus 3 -> new text
        (
            "Proses pengembangan dan implementasi media berdasarkan tahapan model ASSURE.",
            "Efektivitas media web microlearning berbasis Teknik Feynman dalam meningkatkan kemampuan metakognitif dan keterampilan berbicara mahasiswa.",
        ),
        # 5. Subfokus 4 -> DELETE
        (
            "Kelayakan dan efektivitas media berdasarkan penilaian ahli, respons pengguna, dan hasil uji lapangan.",
            "",  # delete
        ),
    ]

    print("=" * 60)
    print("FOKUS PENELITIAN - CORRECTIONS (Round 2b)")
    print("=" * 60)

    print(f"\n📂 Loading document...")
    doc = Document(docx_path)
    print(f"   Paragraphs: {len(doc.paragraphs)}")

    print(f"\n🔄 Performing corrections...")
    total = 0

    for i, (before, after) in enumerate(replacements, 1):
        is_deletion = (after == "")
        found = False

        if is_deletion:
            # Find and remove entire paragraph
            for p in doc.paragraphs:
                if before in p.text:
                    if p.text.strip() == before.strip() or before in p.text:
                        remove_paragraph(p)
                        found = True
                        total += 1
                        break
        else:
            for p in doc.paragraphs:
                if before in p.text:
                    if replace_in_paragraph(p, before, after):
                        found = True
                        total += 1
                        break

        action = "🗑️  DELETE" if is_deletion else "✏️  REPLACE"
        status = "✅" if found else "⚠️  NOT FOUND"
        short = before[:80] + "..." if len(before) > 80 else before
        print(f"   [{i}/{len(replacements)}] {action} {status}")
        print(f"       {short}")

    print(f"\n💾 Saving document (overwriting)...")
    doc.save(docx_path)
    print(f"   Saved to: {docx_path}")

    print(f"\n{'=' * 60}")
    print(f"   Total changes: {total}/{len(replacements)}")
    print(f"{'=' * 60}")


if __name__ == '__main__':
    main()
